"""
QA AI Studio
Production Text Extractor
Version: 2.0
"""

from pathlib import Path

from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from Core.logger import Logger


class TextExtractor:

    def __init__(self):

        self.logger = Logger.get_logger()

    # --------------------------------------------------
    # Extract Text
    # --------------------------------------------------

    def extract_text(self, file_path):

        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(file_path)

        extension = file_path.suffix.lower()

        self.logger.info(
            f"Extracting text from: {file_path.name}"
        )

        if extension == ".pdf":

            text = self._extract_pdf(file_path)

        elif extension == ".docx":

            text = self._extract_docx(file_path)

        elif extension == ".txt":

            text = self._extract_txt(file_path)

        elif extension == ".xlsx":

            text = self._extract_xlsx(file_path)

        else:

            raise ValueError(
                f"Unsupported file type: {extension}"
            )

        text = text.strip()

        self.logger.info(
            f"Extracted {len(text):,} characters."
        )

        return text

    # --------------------------------------------------
    # PDF
    # --------------------------------------------------

    def _extract_pdf(self, file_path):

        reader = PdfReader(file_path)

        pages = []

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:

                pages.append(page_text)

        return "\n\n".join(pages)

    # --------------------------------------------------
    # DOCX
    # --------------------------------------------------

    def _extract_docx(self, file_path):

        document = Document(file_path)

        parts = []

        # ------------------------
        # Paragraphs
        # ------------------------

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                parts.append(text)

        # ------------------------
        # Tables
        # ------------------------

        for table in document.tables:

            for row in table.rows:

                cells = []

                for cell in row.cells:

                    value = cell.text.strip()

                    if value:

                        cells.append(value)

                if cells:

                    parts.append(" | ".join(cells))

        # ------------------------
        # Headers
        # ------------------------

        for section in document.sections:

            for paragraph in section.header.paragraphs:

                text = paragraph.text.strip()

                if text:

                    parts.append(text)

        # ------------------------
        # Footers
        # ------------------------

        for section in document.sections:

            for paragraph in section.footer.paragraphs:

                text = paragraph.text.strip()

                if text:

                    parts.append(text)

        return "\n".join(parts)

    # --------------------------------------------------
    # TXT
    # --------------------------------------------------

    def _extract_txt(self, file_path):

        with open(
            file_path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as file:

            return file.read()

    # --------------------------------------------------
    # XLSX
    # --------------------------------------------------

    def _extract_xlsx(self, file_path):

        workbook = load_workbook(
            file_path,
            data_only=True
        )

        lines = []

        for sheet in workbook.worksheets:

            lines.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):

                values = []

                for value in row:

                    if value is not None:

                        values.append(str(value))

                if values:

                    lines.append(" | ".join(values))

        return "\n".join(lines)