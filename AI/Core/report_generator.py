import os
import shutil
from openpyxl import load_workbook


# ---------------------------------
# Paths
# ---------------------------------
TEMPLATE_FILE = os.path.join(
    "AI",
    "Templates",
    "PSW_Test_Case_Template.xlsx"
)

OUTPUT_FOLDER = os.path.join(
    "AI",
    "Output"
)


# ---------------------------------
# Create Excel Test Case Report
# ---------------------------------
def create_test_case_report(test_cases):

    # Create Output folder if not exists
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)

    output_file = os.path.join(
        OUTPUT_FOLDER,
        "Generated_Test_Cases.xlsx"
    )

    # Copy Template
    shutil.copy(
        TEMPLATE_FILE,
        output_file
    )

    # Open Excel
    workbook = load_workbook(output_file)

    worksheet = workbook["Test Cases"]

    start_row = 2

    for index, tc in enumerate(test_cases):

        row = start_row + index

        worksheet.cell(row=row, column=1).value = index + 1
        worksheet.cell(row=row, column=2).value = tc["scenario"]
        worksheet.cell(row=row, column=3).value = tc["priority"]
        worksheet.cell(row=row, column=4).value = tc["test_case"]
        worksheet.cell(row=row, column=5).value = tc["pre_conditions"]
        worksheet.cell(row=row, column=6).value = tc["steps"]
        worksheet.cell(row=row, column=7).value = tc["expected_result"]
        worksheet.cell(row=row, column=8).value = ""
        worksheet.cell(row=row, column=9).value = "Not Executed"
        worksheet.cell(row=row, column=10).value = tc["test_type"]

    workbook.save(output_file)

    print(f"\nReport Created Successfully\n{output_file}")

    return output_file