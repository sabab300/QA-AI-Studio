# Create: AI/Core/word_exporter.py

"""
QA AI Studio
Word Exporter

Version: 1.0

Exports generated test cases to a formatted Word (.docx) document.
Uses python-docx, which is already a project dependency (also used
by TextExtractor to read .docx knowledge files).

Row schema matches ExcelExporter.parse_llm_output() output:
    scenario, importance, test_type, test_case,
    pre_conditions, steps, expected_result
"""

from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from Core.logger import Logger


class WordExporter:

    def __init__(self):

        self.logger = Logger.get_logger()

    # --------------------------------------------------

    def export_test_cases(

        self,

        rows,

        output_file,

        title="QA AI Studio — Generated Test Cases"

    ):

        document = Document()

        heading = document.add_heading(
            title,
            level=1
        )

        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

        document.add_paragraph(
            f"Total Test Cases: {len(rows)}"
        )

        headers = [
            "S.No",
            "Name / Scenario / Requirement",
            "Importance",
            "Test Type",
            "Test Case",
            "Pre-Conditions",
            "Steps",
            "Expected Result",
            "Actual Result",
        ]

        table = document.add_table(
            rows=1,
            cols=len(headers)
        )

        table.style = "Light Grid Accent 1"

        header_cells = table.rows[0].cells

        for index, header in enumerate(headers):

            header_cells[index].text = header

            for paragraph in header_cells[index].paragraphs:

                for run in paragraph.runs:

                    run.font.bold = True

                    run.font.size = Pt(10)

        for index, row in enumerate(rows, start=1):

            cells = table.add_row().cells

            values = [
                str(index),
                row.get("scenario", ""),
                row.get("importance", ""),
                row.get("test_type", ""),
                row.get("test_case", ""),
                row.get("pre_conditions", ""),
                row.get("steps", ""),
                row.get("expected_result", ""),
                "",
            ]

            for col_index, value in enumerate(values):

                cells[col_index].text = value

                for paragraph in cells[col_index].paragraphs:

                    for run in paragraph.runs:

                        run.font.size = Pt(9)

        Path(output_file).parent.mkdir(
            parents=True,
            exist_ok=True
        )

        document.save(output_file)

        self.logger.info(
            f"Word exported: {output_file}"
        )

        return output_file